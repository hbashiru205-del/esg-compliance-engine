"""Build a Word document of question/answer pairs (python-docx)."""
import io
import re
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor

_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text: str) -> str:
    return _BAD.sub("", text or "")


def _runs(par, text: str):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            par.add_run(part).bold = (i % 2 == 1)


def build_docx(pairs, doc_names=None, title="Compliance Q&A Export",
               brand="Clarix Intelligence") -> bytes:
    d = Document()
    d.core_properties.title = title
    d.core_properties.author = brand
    d.add_heading(_clean(title), level=0)
    meta = d.add_paragraph()
    r = meta.add_run(f"{brand}  |  Generated "
                     f"{datetime.now():%d %b %Y, %H:%M}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    if doc_names:
        m = d.add_paragraph()
        r = m.add_run("Documents indexed: " + _clean(", ".join(doc_names)))
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    for n, p in enumerate(pairs, 1):
        d.add_heading(f"Q{n}. " + _clean(p["question"]), level=2)
        for line in _clean(p["answer"]).split("\n"):
            if not line.strip():
                continue
            m = re.match(r"^\s*[\*\-]\s+(.*)", line)
            par = d.add_paragraph(style="List Bullet" if m else None)
            _runs(par, m.group(1) if m else line.strip())
        if p.get("citations"):
            par = d.add_paragraph()
            r = par.add_run("Cited sources: "
                            + _clean("; ".join(p["citations"])))
            r.italic = True
            r.font.size = Pt(9)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
